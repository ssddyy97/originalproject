"""Editable PPTX and matching HTML preview, based on the user's local sources."""
from pathlib import Path
import sys, html, json, math
ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / '.presentation_tools'))
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement
from PIL import Image

OUT = ROOT / 'original'
ASSET = Path(__file__).resolve().parent
PPT = OUT / 'オリジナル作品発表会_完成版.pptx'
NAVY='111F35'; INK='182A40'; MINT='70E2C0'; TEAL='087F74'; BLUE='3569EA'
WHITE='FFFFFF'; BG='F5F7FA'; MUTED='62748A'; LINE='DCE4EC'; PALE='E5F5EF'; DARKCARD='1C304A'; CORAL='EBA881'
prs=Presentation(); prs.slide_width=Inches(13.333333); prs.slide_height=Inches(7.5)
prs.core_properties.title='オリジナル作品発表会 | ソ・ドンユン | Reserve Hub'
prs.core_properties.author='ソ・ドンユン'
prs.core_properties.subject='自己紹介・キャリアビジョン・Reserve Hub'
slides=[]; current=None
def rgb(c): return RGBColor.from_string(c)
def u(v): return Inches(v/96)
def rect(x,y,w,h,fill,r=0,line=None):
    sh=current['slide'].shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if r else MSO_SHAPE.RECTANGLE,u(x),u(y),u(w),u(h))
    if r: sh.adjustments[0]=0.12
    sh.fill.solid(); sh.fill.fore_color.rgb=rgb(fill)
    if line: sh.line.color.rgb=rgb(line); sh.line.width=Pt(.7)
    else: sh.line.fill.background()
    current['items'].append(dict(kind='rect',x=x,y=y,w=w,h=h,fill=fill,r=r,line=line))
    return sh
def txt(text,x,y,w,h,size=24,color=INK,bold=False,font='Yu Gothic',align='left'):
    sh=current['slide'].shapes.add_textbox(u(x),u(y),u(w),u(h)); tf=sh.text_frame
    tf.clear(); tf.word_wrap=True
    tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
    tf.vertical_anchor=MSO_ANCHOR.TOP
    for i,line in enumerate(text.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.text=line
        p.alignment={'left':PP_ALIGN.LEFT,'center':PP_ALIGN.CENTER,'right':PP_ALIGN.RIGHT}[align]
        p.space_after=Pt(0); p.space_before=Pt(0); p.line_spacing=1.15
        for run in p.runs:
            run.font.name=font; run.font.size=Pt(size*.75); run.font.bold=bold; run.font.color.rgb=rgb(color)
            rp=run._r.get_or_add_rPr(); ea=OxmlElement('a:ea'); ea.set('typeface',font); rp.append(ea)
    current['items'].append(dict(kind='text',text=text,x=x,y=y,w=w,h=h,size=size,color=color,bold=bold,font=font,align=align))
    return sh
def picture(name,x,y,w,h):
    p=ASSET/name; current['slide'].shapes.add_picture(str(p),u(x),u(y),width=u(w),height=u(h))
    current['items'].append(dict(kind='image',src='presentation_assets/'+name,x=x,y=y,w=w,h=h))
def new(section,title=None,sub=None,dark=False,notes='',source=''):
    global current
    s=prs.slides.add_slide(prs.slide_layouts[6]); s.background.fill.solid(); s.background.fill.fore_color.rgb=rgb(NAVY if dark else BG)
    current=dict(slide=s,items=[],dark=dark,title=title or section,notes=notes,source=source); slides.append(current)
    fg=WHITE if dark else INK
    rect(54,40,28,4,MINT if dark else TEAL)
    txt(section.upper(),94,29,1020,26,14,MINT if dark else TEAL,True,font='Aptos')
    if title: txt(title,54,90,1175,65,39,fg,True)
    if sub: txt(sub,56,160,1150,56,20,'BBC9D8' if dark else MUTED)
    rect(54,666,1172,1,'304157' if dark else LINE)
    txt('SO DONGYUN  /  ORIGINAL WORK PRESENTATION',54,681,960,20,11,'A2B2C5' if dark else MUTED,font='Aptos')
    txt(f'{len(slides):02}',1160,676,64,24,16,MINT if dark else TEAL,True,font='Aptos',align='right')
    s.notes_slide.notes_text_frame.text=notes+'\n\n参照：'+source
    return s
def tag(label,x,y,w=180,dark=False):
    rect(x,y,w,32,DARKCARD if dark else PALE,12)
    txt(label,x+12,y+6,w-24,23,14,MINT if dark else TEAL,True)
def card(x,y,w,h,num,title,body,dark=False):
    rect(x,y,w,h,DARKCARD if dark else WHITE,14)
    txt(num,x+24,y+20,w-48,46,34,MINT if dark else TEAL,True,font='Aptos')
    txt(title,x+24,y+81,w-48,72,26,WHITE if dark else INK,True)
    txt(body,x+24,y+166,w-48,h-178,20,'BBC9D8' if dark else MUTED)
def band(text,y=593,dark=False):
    rect(54,y,1172,48,DARKCARD if dark else PALE,8)
    txt(text,76,y+10,1128,34,21,MINT if dark else TEAL,True)

# 01 — Opening
new('Original work presentation',dark=True,
    notes='【目安 15秒】\n本日は、自己紹介と今後の目標、そして制作した予約システム「Reserve Hub」について発表します。\n한국어: 자기소개와 목표, 작품을 소개하며 시작합니다.',source='原本スライド1・README')
txt('オリジナル作品\n発表会',54,145,780,185,66,WHITE,True)
txt('学びを、動くシステムへ。',58,352,780,55,31,MINT,True)
txt('徐東潤  /  ソ・ドンユン',58,503,680,43,28,WHITE,True)
txt('配属先：名古屋',58,556,500,31,20,'BBC9D8')
rect(890,130,336,462,DARKCARD,18)
txt('RESERVE\nHUB',921,171,268,123,51,WHITE,True,font='Aptos')
txt('設計から、運用まで。',921,317,268,36,22,MINT,True)
for i in range(3):
    for j in range(3):
        rect(923+j*82,395+i*44,65,29,MINT if (i,j)==(1,1) else '344A63',6)
txt('JAVA  /  SPRING BOOT  /  POSTGRESQL',921,550,290,24,10,'BBC9D8',font='Aptos')

# 02 — Agenda
new('Contents','今日、お伝えしたいこと','「どんな人か」から、「何をつくったか」へ。',notes='【目安 10秒】\n最初に私自身と日本に来たきっかけをお話しします。その後、強みと3年後の目標、作品の工夫、最後にデモをご紹介します。',source='原本スライド2')
items=[('01','私について','自己紹介・日本に来たきっかけ'),('02','強みと目標','自己PR・3年後の自分'),('03','Reserve Hub','作品概要・設計・テスト・運用'),('04','デモとこれから','予約の流れ・今後の改善')]
for i,(n,t,b) in enumerate(items):
    y=239+i*94; txt(n,62,y,86,60,36,TEAL,True,font='Aptos');txt(t,180,y+1,410,47,29,INK,True);txt(b,630,y+7,590,46,22,MUTED)
    if i<3: rect(180,y+72,1040,1,LINE)

# 03 — Profile
new('01 / About me','経営学から、ITのものづくりへ。','新しい知識を学び、実際に形にすることが好きです。',notes='【目安 35秒】\n韓国出身のソ・ドンユンです。大学では経営学を専攻しました。その後、AI金融ソフトウェア学科やSoldeskでITを学び、Web・モバイルアプリ、機械学習、サーバー構築などに取り組んできました。趣味は旅行、ゲーム、ドラマ、映画、読書、野球です。',source='原本スライド4')
rect(54,236,355,388,NAVY,16)
txt('徐東潤',82,278,290,63,46,WHITE,True)
txt('ソ・ドンユン',84,353,285,39,27,MINT,True)
txt('韓国出身\n名古屋配属',84,435,280,84,24,WHITE)
txt('好奇心を、行動につなげる。',84,569,295,29,17,'BBC9D8')
txt('LEARNING',454,243,700,30,14,TEAL,True,font='Aptos')
txt('大学：経営学専攻',454,289,700,40,29,INK,True)
txt('AI金融ソフトウェア学科  /  Soldesk',454,340,744,36,24,MUTED)
rect(454,399,742,1,LINE)
txt('EXPERIENCE',454,424,710,27,14,TEAL,True,font='Aptos')
txt('Web・モバイルアプリ / 深層学習・機械学習\nサーバー環境構築',454,468,740,86,23,INK)
txt('趣味：旅行・ゲーム・ドラマ・映画・読書・野球',454,584,748,32,19,MUTED)

# 04 — Japan
new('01 / My journey','「好き」が、日本で働くという決意に。','文化への関心を、学習と交流で一歩ずつ行動に変えました。',notes='【目安 40秒】\n子どもの頃から日本のアニメや漫画が好きでした。何度も旅行する中で、日本の文化や生活が自分に合っていると感じました。言語交換アプリで知り合った友人と約1年半交流し、日本語の学習を続けてJLPT N2・N1を取得しました。その経験から、日本で生活し、働くことを決めました。',source='原本スライド5')
journey=[('01','文化への興味','アニメ・漫画を通じて\n日本に興味を持つ'),('02','旅行での実感','何度も日本を訪れ\n文化や生活に親しむ'),('03','交流と学習','約1年半の日本語交流\nJLPT N2・N1を取得'),('04','日本で働く','日本で生活し\n働くことを決意')]
for i,(n,t,b) in enumerate(journey):
    x=54+i*299; rect(x,271,275,260,WHITE,12);txt(n,x+22,290,220,58,43,TEAL,True,font='Aptos');txt(t,x+22,363,237,43,26,INK,True);txt(b,x+22,432,239,72,20,MUTED)
    if i<3: txt('→',x+277,377,24,32,22,TEAL,True)
band('興味を持つ → 続ける → 行動する。この姿勢を、仕事にも。')

# 05 — PR
new('02 / My strengths','学び続け、仕組みで考え、最後まで動かす。','今回の開発で表れた、私の3つの強み。',notes='【目安 35秒】\n私の強みは、学び続けること、仕組みで考えること、最後まで動かすことです。今回の作品では、設計を学びながらドメインを分離し、重複予約をアプリケーションとDBの両方で守り、サーバーで動かすところまで取り組みました。エラー時にもログを確認して原因を追うことを大切にしました。\n한국어: 막연한 장점 대신 이번 작품에서 실제로 한 일을 근거로 설명합니다.',source='原本スライド4・README・配備記録から自己PRを構成')
card(54,235,375,321,'01','学び続ける','未経験の技術を学び\n設計・実装で確かめる。\n\nヘキサゴナル構成を実践')
card(452,235,375,321,'02','仕組みで考える','「動く」だけでなく\n競合時の整合性まで考える。\n\nアプリ＋DBで重複を防止')
card(851,235,375,321,'03','最後まで動かす','テスト・配備・障害分析まで\n一連の流れを経験する。\n\nログを根拠に原因を特定')
band('実装だけで終わらず、動く理由・失敗する理由まで理解する。')

# 06 — Vision
new('02 / Career vision','3年後：周囲から信頼されるテクニカルリードへ。','オープンウインドウ64から、目指す姿と8つの軸を整理。',notes='【目安 35秒】\n3年後は、技術だけでなく周囲から信頼されるテクニカルリードを目指しています。技術力、対話力、リーダーシップを伸ばしながら、語学、健康、メンタル、創造力も大切にします。日々の学習や日本語での交流を続け、小さな行動を積み重ねたいです。\n한국어: 본문은 원본 64표의 중심 8축 요약입니다. 상세 원본은 부록에 보존했습니다.',source='原本スライド8のオープンウインドウ64')
cells=[('技術力','新技術の学習・AWS資格'),('コミュニケーション','傾聴・約束を守る'),('高い成果・収入','仕事の楽しさ・自己成長'),('リーダーシップ','責任感・全体を見る力'),('テクニカルリード','信頼される存在へ'),('メンタル','前向きな思考・感情管理'),('体力','週3回の運動・睡眠'),('言語能力','日本語・IT用語・英語'),('創造力','「なぜ？」・新しい方法')]
for i,(t,b) in enumerate(cells):
    x=54+(i%3)*399;y=232+(i//3)*132;center=i==4
    rect(x,y,374,114,NAVY if center else WHITE,12)
    txt(t,x+20,y+21,336,38,25,MINT if center else INK,True,align='center')
    txt(b,x+16,y+70,342,29,18,WHITE if center else MUTED,align='center')

# 07 — Project chapter
new('03 / Original work',dark=True,notes='【目安 20秒】\nここからは、制作したReserve Hubをご紹介します。会議室などの共有リソースを時間単位で予約するシステムです。安心して予約でき、変更にも対応しやすい基盤を目指しました。',source='README・原本スライド9')
txt('Reserve Hub',54,131,1130,118,91,WHITE,True,font='Aptos')
txt('予約に、確かな仕組みを。',59,290,1120,72,46,MINT,True)
txt('会議室などの共有リソースを管理する、予約プラットフォーム。',60,389,1110,46,25,'BBC9D8')
for i,(a,b) in enumerate([('01','重複を防ぐ'),('02','変更に備える'),('03','運用までつなぐ')]):
    x=58+i*399;rect(x,516,374,87,DARKCARD,12);txt(a,x+20,537,60,40,26,MINT,True,font='Aptos');txt(b,x+93,537,260,40,25,WHITE,True)

# 08 — Why
new('03 / Problem & approach','予約システムは、「登録できる」だけでは足りない。','この題材を通じて、整合性・変更への対応・運用まで取り組みました。',notes='【目安 30秒】\n予約システムを題材にすると、登録や照会に加え、同時アクセス時の競合や状態管理まで考える必要があります。今回は、同じ時間帯に予約が重なる問題を中心に、設計、テスト、デプロイまで経験できる作品にしました。',source='READMEの重要課題・設計選択')
for i,(a,b,c) in enumerate([('同時に予約されたら？','重複予約のリスク','アプリの事前確認＋DB制約'),('技術が変わったら？','予約ルールへの影響','ドメインと外部技術を分離'),('修正を公開するには？','品質確認と配備の手間','テスト・ビルド・配備を自動化')]):
    y=237+i*122;rect(54,y,1172,102,WHITE,12);txt(a,78,y+24,368,44,26,INK,True);txt(b,455,y+18,305,30,18,MUTED);txt(c,455,y+53,707,34,24,TEAL,True)

# 09 — UI
new('03 / Product experience','空き時間の確認から予約まで、ひとつの画面で。','会議室を選ぶ → 時間を選ぶ → 内容を確認する。',notes='【目安 30秒】\n会議室A・B・Cと日付を選び、空き時間を確認できます。緑は空き、青は選択中です。右側で内容を確認して予約します。作成後はDBからの再照会やキャンセルもできます。\n注：この画像は実際のindex.htmlを表示し、空き状況APIにサンプル応答を使用した画面例です。実DBの動作証明ではありません。',source='src/main/resources/static/index.html・画面表示用サンプルデータ')
im=Image.open(ASSET/'reserve_hub_ui.png'); im.crop((95,308,1340,1000)).save(ASSET/'ui_crop.png')
picture('ui_crop.png',54,239,852,474)
# Keep the screenshot inside the content area: resize preserving aspect ratio.
last=current['slide'].shapes[-1];last.height=u(391);last.width=u(704)
current['items'][-1].update(w=704,h=391)
for i,(a,b) in enumerate([('01  空き状況を確認','会議室・日付ごとに表示'),('02  選択内容を確認','時間帯と予約者を一覧化'),('03  予約後も管理','再照会・キャンセルに対応')]):
    y=256+i*118;txt(a,806,y,411,37,25,INK,True);txt(b,806,y+48,411,35,21,MUTED)
txt('実装画面の表示例（空き状況はサンプルデータ）',60,637,1000,22,12,MUTED)

# 10 — concurrency
new('03 / Core design','重複予約は、アプリとDBの二段階で防ぐ。','事前確認の直後に別のリクエストが入っても、最後はDBが整合性を守ります。',notes='【目安 45秒】\nアプリケーションでは、同じリソースの確定済み予約と時間が重なるかを確認します。ただし、複数のリクエストが同時に確認を通る可能性があります。そのため、PostgreSQLの排他制約でも重複を防止しました。時間帯は開始を含み終了を含まないため、10時から11時と11時から12時の連続予約は可能です。キャンセル済みの予約は重複判定から除外します。',source='CreateReservationService.java・V1__create_reservations.sql')
rect(54,238,475,334,WHITE,12)
txt('同じ会議室 / 同じ時間帯',78,265,425,36,25,INK,True)
txt('10:00',159,330,150,27,18,MUTED);txt('11:00',384,330,110,27,18,MUTED)
txt('要求 A',78,385,100,30,20,INK,True);rect(182,378,297,45,TEAL,8);txt('10:00 — 11:00',210,386,250,32,20,WHITE,True)
txt('要求 B',78,463,100,30,20,INK,True);rect(182,456,297,45,'EBDCD3',8);txt('同じ時間帯は競合',210,464,250,32,20,INK,True)
for y,n,t,b in [(238,'1','Application','確定済み予約の重なりを事前確認'),(416,'2','PostgreSQL','排他制約で、同時実行時も最終保護')]:
    rect(560,y,666,156,NAVY if n=='2' else WHITE,12);txt(n,588,y+26,65,70,52,MINT if n=='2' else TEAL,True,font='Aptos');txt(t,674,y+26,500,44,31,WHITE if n=='2' else INK,True,font='Aptos');txt(b,674,y+91,510,35,21,WHITE if n=='2' else MUTED)
band('対象は CONFIRMED のみ。キャンセル後は、同じ時間帯を再予約できます。')

# 11 — Architecture
new('03 / Architecture','変わりやすい技術と、変えたくない予約ルールを分ける。','ヘキサゴナルアーキテクチャ / Ports & Adapters',notes='【目安 45秒】\n中心はReservationとReservationPeriodです。その外側にユースケースを置き、Webと永続化をポートとアダプターで分離しました。DomainとApplication ServiceにはSpringやJPAのアノテーションを置かず、設定クラスで組み立てています。図の矢印は処理の流れです。依存関係ではJPA AdapterがRepository Portを実装します。DBを変える場合、PostgreSQL固有の排他制約は再設計が必要です。',source='reservation/domain・application・adapter・ReservationConfiguration.java')
rect(323,246,627,310,PALE,15)
txt('CORE  /  予約のルールとユースケース',346,266,582,29,17,TEAL,True,font='Yu Gothic')
rect(54,336,222,153,WHITE,12);txt('WEB',76,354,179,28,15,TEAL,True,font='Aptos');txt('REST\nController',76,397,180,78,27,INK,True,font='Aptos')
txt('→',281,389,37,39,32,TEAL,True)
rect(345,335,184,154,WHITE,12);txt('INPUT PORT',361,354,155,25,14,TEAL,True,font='Aptos');txt('Use Case',361,399,160,38,26,INK,True,font='Aptos')
txt('→',533,389,37,39,30,TEAL,True)
rect(574,321,329,182,NAVY,12);txt('APPLICATION / DOMAIN',594,342,290,27,14,MINT,True,font='Aptos');txt('Reservation',594,389,287,44,32,WHITE,True,font='Aptos');txt('状態と時間のルール',594,447,280,30,21,WHITE)
txt('Repository Port',604,522,288,25,17,TEAL,True,font='Aptos')
txt('→',962,389,36,39,30,TEAL,True)
rect(1007,336,219,153,WHITE,12);txt('PERSISTENCE',1026,354,185,28,14,TEAL,True,font='Aptos');txt('JPA Adapter\nPostgreSQL',1026,401,190,75,24,INK,True,font='Aptos')
band('中心のルールは、Spring MVCやJPAに直接依存しない。')
txt('矢印：処理の流れを簡略化',54,642,1000,19,12,MUTED)

# 12 — testing
new('03 / Verification','「同時に来ても1件だけ」を、テストで確かめる。','APIレベルの同時実行テストで、レスポンスと保存件数を確認。',notes='【目安 35秒】\n同じ予約を10件同時に送信し、成功が1件、競合が9件、DBに保存された予約が1件になることを検証しています。既存のローカルJUnitレポートでは29件のテストが成功しています。これは保存済みの実行結果で、この資料作成時に新しく実行した結果ではありません。\n한국어: 수치는 저장된 test-results와 테스트 코드에서 확인한 값입니다.',source='ReservationApiConcurrencyTest.java・build/test-results/test/TEST-*.xml（既存レポート）')
for x,num,label,col in [(54,'10','同時リクエスト',INK),(453,'1','201 Created',TEAL),(852,'9','409 Conflict',INK)]:
    rect(x,244,374,258,WHITE,12);txt(num,x+24,263,326,140,104,col,True,font='Aptos',align='center');txt(label,x+22,425,330,42,25,MUTED,True,align='center')
rect(54,526,1172,107,NAVY,12);txt('DB保存：1件',79,549,380,48,33,MINT,True);txt('既存JUnitレポート：29件成功 / 失敗・エラー 0件\nDomain・Application・Persistence・Webを検証',485,547,711,70,20,WHITE)

# 13 — Delivery
new('03 / Delivery','テスト済みの成果物を、そのままサーバーへ。','GitHub Actionsで、検証から配備までをつなぐ。',notes='【目安 35秒】\nmainブランチにPushすると、GitHub上でテストとJARのビルドを実行します。テストが成功した場合だけ、Oracle LinuxのSelf-hosted Runnerが成果物を受け取り、配備スクリプトを実行します。運用環境はNginx、systemd、Docker上のPostgreSQLを利用しています。配備記録では、既存JARのバックアップと起動失敗時のロールバックも構成しました。',source='.github/workflows/ci.yml・Reserve-Hub_CICD_Deployment_Setup.md')
steps=[('01','Push','main branch'),('02','Test & Build','GitHub-hosted'),('03','JAR Artifact','検証済み成果物'),('04','Deploy','Self-hosted')]
for i,(n,t,b) in enumerate(steps):
    x=54+i*299;rect(x,257,275,200,NAVY if i==3 else WHITE,12);txt(n,x+23,277,225,39,27,MINT if i==3 else TEAL,True,font='Aptos');txt(t,x+23,333,234,43,29,WHITE if i==3 else INK,True,font='Aptos');txt(b,x+23,399,233,31,18,'BBC9D8' if i==3 else MUTED)
    if i<3:txt('→',x+277,338,25,35,24,TEAL,True)
rect(54,500,1172,132,PALE,12);txt('ORACLE LINUX',77,518,1050,27,15,TEAL,True,font='Aptos')
txt('Nginx  →  Spring Boot / systemd  →  PostgreSQL / Docker',77,558,1115,42,29,INK,True,font='Aptos')

# 14 — Troubleshooting
new('03 / Troubleshooting','エラーの表面ではなく、ログから原因を探る。','配備環境で直面した3つの問題と、その解決。',notes='【目安 45秒】\nRunnerの認証失敗は、VMの時刻が約15時間50分ずれていたことが原因でした。chronyで同期して解決しました。systemdの203/EXECはSELinuxのログを調べ、Runnerの配置場所とコンテキストを修正しました。Nginxの502も監査ログから通信制限を確認しました。SELinux全体を無効化せず、必要な設定を修正して解決しました。',source='READMEのトラブルシューティング・配備記録')
for x,t in [(76,'発生した問題'),(427,'確認した原因'),(797,'対応')]:txt(t,x,231,385,33,18,TEAL,True)
rows=[('Runner 認証失敗','VMの時刻ずれ','chronyで時刻を同期'),('systemd 203/EXEC','SELinuxによる実行制限','/optへ移動・コンテキスト復元'),('Nginx 502','SELinuxによる通信制限','必要なネットワーク設定を許可')]
for i,(a,b,c) in enumerate(rows):
    y=278+i*93;rect(54,y,1172,78,WHITE,9);txt(a,77,y+22,345,38,23,INK,True);txt(b,427,y+23,355,37,21,MUTED);txt(c,797,y+24,405,37,20,INK)
band('大切にしたこと：ログで原因を特定し、必要な変更だけを加える。')

# 15 — Demo
new('04 / Live demo','予約 → 再照会 → キャンセル → 再予約', '画面の変化と、保存された状態をセットで確認します。',dark=True,notes='【目安 60〜90秒】\nそれでは実際の画面をご覧ください。会議室と空いている時間帯を選び、予約します。CONFIRMEDになったことを確認し、DBから再照会します。次にキャンセルするとCANCELLEDに変わり、時間帯が空きに戻ります。同じ時間帯を再予約できることを確認します。\n한국어: 브라우저에서 실제 서버를 열어 진행하세요. 발표 전 DB와 서버 실행 및 사용할 시간의 공석을 확인하세요. 이 PPT는 DB나 서버를 시작하지 않습니다.',source='index.html・既存発表スクリプト')
for i,(n,t,b) in enumerate([('01','予約する','CONFIRMED\n空き → 予約済み'),('02','DBから再照会','保存された予約を\n再取得'),('03','キャンセル','CANCELLED\n予約済み → 空き'),('04','同じ時間に再予約','履歴を残しながら\n時間帯を再利用')]):
    x=54+i*299;rect(x,260,275,292,DARKCARD,12);txt(n,x+22,284,230,52,38,MINT,True,font='Aptos');txt(t,x+22,363,235,72,25,WHITE,True);txt(b,x+22,461,235,74,20,'BBC9D8')
band('確認ポイント：画面の表示だけでなく、DBに保存された状態を見る。',dark=True)

# 16 — Next
new('04 / Next steps','次は、デモから実運用へ。','現在の制約を整理し、改善の優先順位を明確にする。',notes='【目安 30秒】\n現在は各時間帯ごとに空き状況APIを呼んでいます。最初に日単位の一括取得APIへ改善したいです。次に認証・認可、ユーザー別予約管理、会議室情報のDB管理を追加し、HTTPSや監視も整えます。これらは今後の改善予定であり、実装済みではありません。',source='READMEの設計選択・今後の改善計画')
card(54,239,375,332,'NEXT 01','一覧取得を効率化','時間帯ごとのAPI呼び出し\n\n→ 日付・会議室単位の\n   一括取得APIへ')
card(452,239,375,332,'NEXT 02','利用者ごとの管理','認証・認可を追加\n\n→ 自分の予約一覧\n→ 会議室情報のDB管理')
card(851,239,375,332,'NEXT 03','運用の見える化','HTTPS・監視を整備\n\n→ Secret管理を強化\n→ 障害の早期発見')
txt('上記は今後の改善予定です。',58,607,1120,32,18,MUTED)

# 17 — Close
new('Thank you',dark=True,notes='【目安 20秒】\nこの作品を通じて、実装だけでなく設計、テスト、デプロイ、障害分析まで一つの流れとして経験できました。今後も学び続け、安心して任せてもらえるエンジニアを目指します。ご清聴ありがとうございました。',source='既存発表スクリプト・原本最終スライドの本人情報に修正')
txt('つくる。確かめる。\n動かし続ける。',54,153,1172,179,63,WHITE,True)
txt('学びを積み重ね、信頼されるエンジニアへ。',58,383,1160,58,32,MINT,True)
txt('ご清聴ありがとうございました。',59,500,1155,43,27,WHITE)
txt('徐東潤  /  ソ・ドンユン   ｜   名古屋',59,571,1130,34,22,'BBC9D8')

# 18 — original map (appendix)
new('Appendix A / Original vision map','オープンウインドウ64：原本','本人が作成した、3年後の目標と行動計画。',notes='【補足資料 / 通常は読み上げ不要】\n原本スライド8の画像をそのまま保存しています。本文スライド6は中心目標と8つの軸を読みやすく要約したものです。',source='原本スライド8 / ppt/media/image11.png')
im=Image.open(ASSET/'image11.png'); ratio=im.width/im.height; h=418;w=h*ratio;picture('image11.png',54,227,w,h)
txt('中心の目標',751,260,420,33,19,TEAL,True)
txt('テクニカル\nリード',751,308,420,115,42,INK,True)
txt('技術・対話・健康・語学。\n継続できる日々の行動へ。',751,480,427,91,23,MUTED)

# 19 — Q&A
new('Appendix B / Design decisions','設計について聞かれたら。','選択理由だけでなく、トレードオフも説明できるように。',notes='【補足資料 / 質疑応答用】\nQ：小規模なのに複雑では？ A：一般的な三層構造の方が実装量は少ないです。今回は設計の学習と変更の影響範囲の明確化を目的に選びました。\nQ：DBは交換できますか？ A：保存処理はAdapterで分離していますが、重複防止制約はPostgreSQL固有なので再設計が必要です。\nQ：キャンセル時に消さない理由は？ A：履歴を残すためです。確定済みだけを重複判定の対象にします。',source='READMEのトレードオフ・既存発表スクリプトの想定Q&A')
qas=[('なぜヘキサゴナル？','予約ルールを外部技術から分離し、テストしやすくする。','代償：小規模でもクラス数が増える。'),('なぜ静的HTML？','予約ロジックと配備までの完成を優先し、配布単位をひとつに。','今後：画面や規模に応じて構成を見直す。'),('DBは交換できる？','Repository Portを境界に、保存の実装を分離。','注意：PostgreSQL固有の排他制約は再設計が必要。')]
for i,(q,a,b) in enumerate(qas):
    y=229+i*137;rect(54,y,1172,116,WHITE,10);txt(q,78,y+29,350,49,25,INK,True);txt(a,442,y+20,756,39,22,INK);txt(b,442,y+69,756,33,18,TEAL)

prs.save(PPT)

# Matching vector/text HTML enables local preview and PDF export without Office.
pages=[]
for idx,s in enumerate(slides):
    els=[]
    for a in s['items']:
        style=f"left:{a['x']}px;top:{a['y']}px;width:{a['w']}px;height:{a['h']}px;"
        if a['kind']=='rect':
            style+=f"background:#{a['fill']};border-radius:{a['r']}px;"
            if a['line']:style+=f"border:1px solid #{a['line']};"
            els.append(f'<div class="shape" style="{style}"></div>')
        elif a['kind']=='image':els.append(f'<img class="shape" style="{style}" src="{a["src"]}">')
        else:
            style+=f"font-size:{a['size']}px;color:#{a['color']};font-weight:{700 if a['bold'] else 400};font-family:'{a['font']}',sans-serif;text-align:{a['align']};"
            els.append(f'<div class="shape text" style="{style}">{html.escape(a["text"])}</div>')
    pages.append(f'<section class="slide" id="slide-{idx+1}" style="background:#{NAVY if s["dark"] else BG}">'+''.join(els)+'</section>')
preview='''<!doctype html><html lang="ja"><meta charset="utf-8"><title>オリジナル作品発表会 — Preview</title><style>
*{box-sizing:border-box}body{margin:0;background:#dce4ec}.slide{position:relative;width:1280px;height:720px;margin:24px auto;overflow:hidden;break-after:page}.shape{position:absolute}.text{white-space:pre-wrap;line-height:1.15;overflow:visible;word-break:normal;overflow-wrap:normal}
@page{size:13.333333in 7.5in;margin:0}@media print{body{background:white}.slide{margin:0;break-after:page;-webkit-print-color-adjust:exact;print-color-adjust:exact}}
</style><body>'''+''.join(pages)+'</body></html>'
(OUT/'オリジナル作品発表会_プレビュー.html').write_text(preview,encoding='utf-8')
notes=['# 완성본 발표 가이드','', '본문 17장 + 부록 2장. 일본어 발표자 노트는 PPT 각 슬라이드에도 들어 있습니다.', '', '전체 발표는 데모 포함 약 9~11분을 기준으로 구성했습니다. 작품 소개만 발표할 때는 7~17번을 사용하세요.', '', '화면 이미지는 실제 프로젝트 HTML을 표시하고 API에 샘플 응답을 넣어 캡처했습니다. 실DB 동작 증빙은 아닙니다. 테스트 수치는 저장된 로컬 JUnit 결과이며 이번 작업에서 재실행한 값은 아닙니다. 원본 PPT와 소스 코드는 수정하지 않았습니다.', '', '3년 후 목표의 본문은 원본 표를 요약한 내용입니다. 자세한 원본은 18번 부록에 보존했습니다. 자기 PR은 실제 프로젝트 경험을 근거로 정리했습니다.']
for i,s in enumerate(slides):notes.extend(['',f'## {i+1:02}. {s["title"]}','',s['notes'],'',f'자료 근거: {s["source"]}'])
(OUT/'完成版_発表ガイド_JP-KR.md').write_text('\n'.join(notes),encoding='utf-8')
(ASSET/'deck_manifest.json').write_text(json.dumps([dict(title=s['title'],items=s['items']) for s in slides],ensure_ascii=False,indent=2),encoding='utf-8')
print(f'Created {PPT.name}: {len(slides)} slides')
